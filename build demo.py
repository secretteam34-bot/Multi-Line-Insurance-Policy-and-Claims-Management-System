import os
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, PageBreak, Table, TableStyle
import imageio
import numpy as np

BASE = r"C:\Users\PRAVEEN PRABAKARN\Desktop\new fold\Multi-Line Insurance Policy and Claims_NM"
SHOT = os.path.join(BASE, "screenshots")
OUT = os.path.join(BASE, "reports")
os.makedirs(OUT, exist_ok=True)

TITLE = "Multi-Line Insurance Policy and Claims Management System"
SUB = "Salesforce Developer - Group Project Report"
AUTHOR = "Praveen Prabakarn"
LINK = ("https://myskillwallet.ai/dashboard/skillwallet/module/salesforce-developer-nm-eng-6a69e3b28beabdd402737070/"
        "group-projects/6a6b1d37dafa0b21ea8a0900/Multi-Line-Insurance-Policy-and-Claims-Management-System-"
        "6ab601b2fa5e5219dc38104e?tab=2")
ORG = "orgfarm-63e3e46d5b-dev-ed (Salesforce Developer Edition)"

# (section, [(file, caption)])
SECTIONS = [
 ("1. Data Model Setup - Policy and Claim Objects", [
  ("01-policy.png", "Creating the custom object Policy (Label: Policy, Plural: Policies, Record Name: Policy Name, Auto Number)."),
  ("02-policy-save.png", "Object creation options and Save for the Policy object."),
  ("03-policy-object.png", "New record type 'Auto' on the Policy object."),
  ("04-policy-auto.png", "Auto record type created and Active (Record Type Name: Auto)."),
  ("05-policy-vehicle.png", "Field Set 'Vehicle' created on Policy - available fields."),
  ("06-policy-fieldsets.png", "Vehicle field set with VIN and Model Year added."),
  ("07-policy-fieldsets.png", "Vehicle field set containing VIN, Model Year, Square Footage, Year Built, Beneficiary Name, Policy Name."),
  ("08-claims.png", "Creating the custom object Claim."),
  ("09-claims-fields.png", "Claim fields: Adjuster, Approval Status, Claim Amount, Date of Loss, Description, Policy lookup, Record Type."),
  ("12-policy-validation-rules.png", "Validation rule VIN_Must_Be_17_Characters on Policy (Auto record type)."),
 ]),
 ("2. AutoQuotingFlow - Screen Flow", [
  ("screenshot-1790923897662-0.jpg", "New Screen Flow canvas (Start and End) in Flow Builder."),
  ("screenshot-1790923981318-1.jpg", "Get Vehicle RT ID - Record Type where SObject Type Name = Policy__c AND Name = Auto."),
  ("screenshot-1790923994351-2.jpg", "Get Vehicle RT ID element added after Start."),
  ("screenshot-1790924247582-3.jpg", "Limitation found: no Policy State picklist field exists, so the Picklist Choice Set could not be built."),
  ("10-flowbuilder.png", "Flow canvas: Get Vehicle RT ID, Basic Information screen, Vehicle-Specific Details screen, Create Draft Policy."),
  ("11-flowbuilder-activation.png", "Create Draft Policy - manual field mapping and varPolicyId storage."),
  ("screenshot-1790926588228-5.jpg", "AutoQuotingFlow V1 saved and Active ('Your automation was activated')."),
 ]),
 ("3. Claims Automation and Environment Checks", [
  ("screenshot-1790933980584-6.jpg", "Apex Classes list - PremiumCalculator class is present in the org."),
  ("screenshot-1790934024002-7.jpg", "Queues page - no queues exist yet (needed for claim routing)."),
  ("screenshot-1790934981959-8.jpg", "Submission Automation Flow (record-triggered on Claim) saved and Active, with the Submit Claim to Approval action."),
 ]),
]

SUMMARY = [
 ("Policy object, Auto record type, Vehicle field set, VIN validation rule", "Completed"),
 ("Claim object with Approval Status and Claim Amount fields", "Completed"),
 ("AutoQuotingFlow (Get Vehicle RT ID, two screens, Create Draft Policy, varPolicyId)", "Completed and Active (Customer, Policy Start Date and Policy State mappings pending missing fields)"),
 ("Submission Automation Flow (Claim over 50000 -> approval)", "Completed and Active (approval process still to be created)"),
 ("Claim Approver Screen Flow", "Partially built, not saved"),
 ("Calculate Premium action / Update Policy With Premium", "Pending - Premium field missing on Policy"),
 ("Claim routing flow, Apex controller and tests, LWC dashboard", "Pending - queues and Policy.Customer__c missing"),
 ("Approval process, sharing rules, permission sets", "Pending - requires administrator access"),
]
ABOUT = ("This report documents the Salesforce configuration carried out for the Multi-Line Insurance Policy and "
         "Claims Management System project, using screenshots captured from the Developer Edition org while building "
         "the data model and the automation. Items that could not be completed because of missing prerequisites or "
         "insufficient privileges are listed honestly in the status table.")

def imgsize(p, maxw):
    im = Image.open(p); w, h = im.size
    return maxw, maxw * h / w

# ---------------- DOCX ----------------
def build_docx():
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.8)
    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(TITLE); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = RGBColor(0x03, 0x2D, 0x60)
    p = d.add_paragraph(SUB); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Prepared by: ").bold = True; p.add_run(AUTHOR)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Salesforce org: ").bold = True; p.add_run(ORG)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Project link: ").bold = True; p.add_run(LINK).font.size = Pt(8)
    d.add_heading("About this report", 1); d.add_paragraph(ABOUT)
    d.add_heading("Work status", 1)
    tb = d.add_table(rows=1, cols=2); tb.style = "Light Grid Accent 1"
    tb.rows[0].cells[0].text = "Item"; tb.rows[0].cells[1].text = "Status"
    for a, b in SUMMARY:
        c = tb.add_row().cells; c[0].text = a; c[1].text = b
    for sec, items in SECTIONS:
        d.add_page_break(); d.add_heading(sec, 1)
        for f, cap in items:
            d.add_picture(os.path.join(SHOT, f), width=Inches(6.4))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c = d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = c.add_run("Figure: " + cap); rr.italic = True; rr.font.size = Pt(9)
    d.save(os.path.join(OUT, "Multi-Line_Insurance_Report.docx"))

# ---------------- PDF ----------------
def build_pdf():
    ss = getSampleStyleSheet()
    ti = ParagraphStyle("t", parent=ss["Title"], textColor=colors.HexColor("#032D60"), fontSize=22)
    cen = ParagraphStyle("c", parent=ss["Normal"], alignment=1)
    cap = ParagraphStyle("cap", parent=ss["Italic"], alignment=1, fontSize=9)
    small = ParagraphStyle("s", parent=cen, fontSize=7)
    doc = SimpleDocTemplate(os.path.join(OUT, "Multi-Line_Insurance_Report.pdf"), pagesize=A4,
                            leftMargin=45, rightMargin=45, topMargin=45, bottomMargin=45)
    st = [Paragraph(TITLE, ti), Paragraph(SUB, cen), Spacer(1, 8),
          Paragraph("<b>Prepared by:</b> " + AUTHOR, cen), Paragraph("<b>Salesforce org:</b> " + ORG, cen),
          Paragraph("<b>Project link:</b> " + LINK, small), Spacer(1, 14),
          Paragraph("About this report", ss["Heading2"]), Paragraph(ABOUT, ss["Normal"]), Spacer(1, 10),
          Paragraph("Work status", ss["Heading2"])]
    rows = [["Item", "Status"]] + [[Paragraph(a, ss["Normal"]), Paragraph(b, ss["Normal"])] for a, b in SUMMARY]
    t = Table(rows, colWidths=[260, 235])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE9F7")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    st.append(t)
    for sec, items in SECTIONS:
        st.append(PageBreak()); st.append(Paragraph(sec, ss["Heading1"]))
        for f, c in items:
            p = os.path.join(SHOT, f); w, h = imgsize(p, 480)
            if h > 330: w, h = w * 330 / h, 330
            st += [RLImage(p, width=w, height=h), Paragraph("Figure: " + c, cap), Spacer(1, 10)]
    doc.build(st)

# ---------------- VIDEO ----------------
def font(sz):
    for n in ("arial.ttf", "segoeui.ttf", "calibri.ttf"):
        try: return ImageFont.truetype(n, sz)
        except Exception: pass
    return ImageFont.load_default()

W, H = 1280, 720
def wrap(draw, text, f, maxw):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=f) <= maxw: cur = t
        else: lines.append(cur); cur = w
    lines.append(cur); return lines

def card(lines, sub=None):
    im = Image.new("RGB", (W, H), (3, 45, 96)); d = ImageDraw.Draw(im)
    y = 230
    for l in lines:
        for x in wrap(d, l, font(46), W - 160):
            d.text((80, y), x, fill="white", font=font(46)); y += 62
    if sub:
        y += 20
        for x in wrap(d, sub, font(26), W - 160):
            d.text((80, y), x, fill=(190, 215, 245), font=font(26)); y += 36
    return im

def slide(path, cap):
    im = Image.new("RGB", (W, H), (245, 247, 250)); d = ImageDraw.Draw(im)
    src = Image.open(path).convert("RGB")
    r = min((W - 60) / src.width, (H - 150) / src.height)
    src = src.resize((int(src.width * r), int(src.height * r)))
    im.paste(src, ((W - src.width) // 2, 20))
    d.rectangle([0, H - 110, W, H], fill=(3, 45, 96))
    y = H - 100
    for x in wrap(d, cap, font(26), W - 80)[:3]:
        d.text((40, y), x, fill="white", font=font(26)); y += 32
    return im

def build_video():
    frames = [(card([TITLE], "Demo walkthrough - " + AUTHOR), 4)]
    for sec, items in SECTIONS:
        frames.append((card([sec]), 2.5))
        for f, c in items:
            frames.append((slide(os.path.join(SHOT, f), c), 5))
    frames.append((card(["Thank you"], "Pending items need admin access: Policy fields, queues, approval process, permission sets."), 5))
    wr = imageio.get_writer(os.path.join(OUT, "Multi-Line_Insurance_Demo.mp4"), fps=10, codec="libx264",
                            quality=7, macro_block_size=16)
    for im, sec in frames:
        a = np.array(im)
        for _ in range(int(sec * 10)): wr.append_data(a)
    wr.close()

build_docx(); build_pdf(); build_video()
print("done")
